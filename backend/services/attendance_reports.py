from collections import defaultdict
from datetime import datetime, time, timedelta, timezone
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

VIETNAM = timezone(timedelta(hours=7))

COPY = {
    "en": {
        "title": "Teacher attendance report", "summary": "Summary", "period": "Period",
        "teacher": "Teacher", "code": "Teacher ID", "confirmed": "Confirmed sessions", "recorded": "Recorded sessions",
        "warning": "Sessions to review", "unknown": "Unverified events", "minutes": "Confirmed teaching minutes",
        "date": "Date", "weekday": "Weekday", "center": "Center", "class": "Class",
        "shift": "Scheduled shift", "in": "Check-in", "out": "Check-out", "result": "Result", "notes": "Notes",
        "CONFIRMED": "Confirmed", "WARNING": "Needs review", "UNVERIFIED": "Unverified",
        "MISSING_CHECK_IN": "Missing check-in", "MISSING_CHECK_OUT": "Missing check-out",
        "LATE": "Late arrival", "LEFT_EARLY": "Early departure", "INVALID_LOCATION": "Invalid location",
        "UNKNOWN_SHIFT": "Historical shift not saved; not counted automatically",
        "UNMATCHED": "No matched class", "INVALID_ORDER": "Checkout precedes check-in",
        "INVALID_SHIFT": "Saved shift has an invalid time range",
        "MANUAL": "Manually recorded", "UNLINKED": "No saved check-in/checkout link",
        "student_count": "Students in visit", "student_names": "Student names (whole visit)",
        "students_unknown": "Not recorded",
        "empty": "No recorded attendance in this period.",
        "policy": "Each saved class counts separately, including consecutive classes sharing one check-in/out. Late/early or incomplete sessions require review and are not included in confirmed totals. No absence is inferred from the current schedule.",
        "days": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"],
    },
    "vi": {
        "title": "Báo cáo chấm công giáo viên", "summary": "Tổng hợp", "period": "Khoảng thời gian",
        "teacher": "Giáo viên", "code": "Mã giáo viên", "confirmed": "Ca đủ điều kiện", "recorded": "Ca đã ghi nhận",
        "warning": "Ca cần xem xét", "unknown": "Sự kiện chưa xác minh", "minutes": "Số phút dạy đủ điều kiện",
        "date": "Ngày", "weekday": "Thứ", "center": "Cơ sở", "class": "Ca dạy",
        "shift": "Giờ ca đã lưu", "in": "Giờ vào", "out": "Giờ ra", "result": "Kết quả", "notes": "Ghi chú",
        "CONFIRMED": "Đủ điều kiện", "WARNING": "Cần xem xét", "UNVERIFIED": "Chưa xác minh",
        "MISSING_CHECK_IN": "Thiếu giờ vào", "MISSING_CHECK_OUT": "Thiếu giờ ra",
        "LATE": "Đi trễ", "LEFT_EARLY": "Về sớm", "INVALID_LOCATION": "Vị trí không hợp lệ",
        "UNKNOWN_SHIFT": "Chưa lưu lịch sử giờ ca; không tự động tính ca",
        "UNMATCHED": "Không khớp ca dạy", "INVALID_ORDER": "Giờ ra trước giờ vào",
        "INVALID_SHIFT": "Khoảng giờ ca đã lưu không hợp lệ",
        "MANUAL": "Ghi nhận thủ công", "UNLINKED": "Chưa lưu liên kết giữa giờ vào và giờ ra",
        "student_count": "Số bé trong buổi", "student_names": "Tên học viên (cả buổi)",
        "students_unknown": "Chưa ghi nhận",
        "empty": "Không có chấm công trong khoảng thời gian này.",
        "policy": "Mỗi ca đã lưu được tính riêng, kể cả các ca liên tục dùng chung một lần vào/ra. Ca trễ, về sớm hoặc thiếu chấm công cần xem xét, chưa cộng vào số ca đủ điều kiện. Không suy ra vắng từ lịch hiện tại.",
        "days": ["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu", "Thứ Bảy", "Chủ Nhật"],
    },
}

COPY["vi"].update({
    "CONFIRMED": "Hợp lệ", "WARNING": "Có bất thường", "INVALID": "Không hợp lệ",
    "UNVERIFIED": "Chưa đủ dữ liệu", "invalid": "Ca không hợp lệ", "OVERTIME": "Ra muộn (overtime)",
    "LEGACY_DOUBLE": "Ca đôi suy luận từ dữ liệu cũ", "LEGACY_SINGLE": "Ca đơn ghép từ dữ liệu cũ",
    "UNKNOWN_SHIFT": "Chưa lưu giờ ca; chưa tính được số phút dạy", "block": "Số ca trong buổi",
    "policy": "Kết quả: Hợp lệ / Có bất thường / Không hợp lệ / Chưa đủ dữ liệu. Mỗi dòng là một buổi điểm danh; cột số ca cho biết buổi đó gồm bao nhiêu ca dạy.",
})
COPY["en"].update({
    "INVALID": "Invalid", "invalid": "Invalid sessions", "OVERTIME": "Overtime",
    "LEGACY_DOUBLE": "Double shift inferred from legacy events", "LEGACY_SINGLE": "Single shift paired from legacy events",
    "UNKNOWN_SHIFT": "Shift times unavailable; teaching minutes unknown", "block": "Sessions in block",
    "policy": "Results: Confirmed / Needs review / Invalid / Unverified. Each row is one attendance visit; the session count shows how many teaching sessions it includes.",
})


def local_datetime(value):
    if isinstance(value, str):
        value = datetime.fromisoformat(value.replace("Z", "+00:00"))
    return value.astimezone(VIETNAM)


def saved_shift(record):
    if not record:
        return []
    if record.get("session_block"):
        return record["session_block"]
    class_id = record.get("session_class_id") or record.get("class_id")
    if not class_id:
        return []
    return [{
        "class_id": class_id, "center_id": record.get("session_center_id"),
        "center_name": record.get("nearest_center_name"),
        "start_time": record.get("session_start_time"), "end_time": record.get("session_end_time"),
    }]


def infer_legacy_blocks(records):
    """Pair only adjacent, unlinked legacy events; never consult today's schedule."""
    records = [dict(record) for record in records]
    groups = defaultdict(list)
    linked = {str(r.get("check_in_record_id")) for r in records if r.get("check_in_record_id")}
    for record in records:
        groups[(record["teacher_id"], local_datetime(record["checked_at"]).date())].append(record)
    for group in groups.values():
        group.sort(key=lambda r: (local_datetime(r["checked_at"]), str(r["record_id"])))
        for entered, exited in zip(group, group[1:]):
            if (entered["event_type"] != "CHECK_IN" or exited["event_type"] != "CHECK_OUT"
                    or entered.get("session_block") or entered.get("session_start_time") or exited.get("check_in_record_id")
                    or str(entered["record_id"]) in linked):
                continue
            center = entered.get("nearest_center_name")
            first, last = entered.get("class_id"), exited.get("class_id")
            if not center or center != exited.get("nearest_center_name") or not first or not last:
                continue
            minutes = (local_datetime(exited["checked_at"]) - local_datetime(entered["checked_at"])).total_seconds() / 60
            double = first != last
            if not 0 < minutes <= 240 or (double and minutes < 180):
                continue
            # A class with its own opposite event is not an orphaned double shift.
            if double and any(
                (r["event_type"] == "CHECK_OUT" and r.get("class_id") == first)
                or (r["event_type"] == "CHECK_IN" and r.get("class_id") == last)
                for r in group
            ):
                continue
            if double and (entered.get("session_start_time") or exited.get("session_start_time")):
                continue
            entered["session_block"] = [
                {"class_id": class_id, "center_name": center, "start_time": None, "end_time": None}
                for class_id in ([first, last] if double else [first])
            ]
            entered["legacy_inferred"] = "LEGACY_DOUBLE" if double else "LEGACY_SINGLE"
            exited["check_in_record_id"] = entered["record_id"]
            linked.add(str(entered["record_id"]))
    return records


def build_attendance_report(teachers, records, start_date, end_date):
    records = infer_legacy_blocks(records)
    checkouts = defaultdict(list)
    legacy = defaultdict(list)
    for record in records:
        if record["event_type"] == "CHECK_OUT" and record.get("check_in_record_id"):
            checkouts[str(record["check_in_record_id"])].append(record)
        elif not record.get("session_block"):
            key = (record["teacher_id"], local_datetime(record["checked_at"]).date(),
                   record.get("session_class_id") or record.get("class_id"))
            legacy[key].append(record)

    used = set()
    details = []

    def append_session(check_in, check_out, linked):
        record = check_in or check_out
        day = local_datetime(record["checked_at"]).date()
        shifts = saved_shift(check_in or check_out)
        if not shifts:
            shifts = [{"class_id": None, "start_time": None, "end_time": None}]
        for shift in shifts:
            warnings = []
            notes = []
            inferred = record.get("legacy_inferred")
            if inferred:
                notes.append(inferred)
            entered = local_datetime(check_in["checked_at"]) if check_in else None
            exited = local_datetime(check_out["checked_at"]) if check_out else None
            start = shift.get("start_time")
            end = shift.get("end_time")
            if not start or not end:
                (notes if inferred else warnings).append("UNKNOWN_SHIFT" if shift["class_id"] else "UNMATCHED")
                minutes = 0
            else:
                start = time.fromisoformat(str(start))
                end = time.fromisoformat(str(end))
                expected_start = datetime.combine(day, start, VIETNAM)
                expected_end = datetime.combine(day, end, VIETNAM)
                minutes = max(0, int((expected_end - expected_start).total_seconds() / 60))
                if expected_end <= expected_start:
                    warnings.append("INVALID_SHIFT")
                if entered and entered > expected_start:
                    warnings.append("LATE")
                if exited and exited < expected_end:
                    warnings.append("LEFT_EARLY")
            if not entered:
                warnings.append("MISSING_CHECK_IN")
            if not exited:
                warnings.append("MISSING_CHECK_OUT")
            if entered and exited and exited < entered:
                warnings.append("INVALID_ORDER")
            if any(event and event["status"] not in ("VALID", "MANUAL") for event in (check_in, check_out)):
                warnings.append("INVALID_LOCATION")
            if any(event and event["status"] == "MANUAL" for event in (check_in, check_out)):
                notes.append("MANUAL")
            if not linked:
                warnings.append("UNLINKED")
            if inferred:
                if check_in.get("attendance_status") == "LATE" and shift == shifts[0]:
                    warnings.append("LATE")
                if check_out and check_out.get("attendance_status") in ("LEFT_EARLY", "OVERTIME") and shift == shifts[-1]:
                    warnings.append(check_out["attendance_status"])
                if check_in.get("attendance_status") not in ("ON_TIME", "EARLY", "LATE", "MANUAL"):
                    warnings.append("UNMATCHED")
            elif exited and end and shift == shifts[-1] and exited > expected_end + timedelta(minutes=30):
                warnings.append("OVERTIME")
            result = ("INVALID" if any(code in warnings for code in ("INVALID_LOCATION", "INVALID_ORDER"))
                      else "UNVERIFIED" if any(code in warnings for code in ("UNKNOWN_SHIFT", "UNMATCHED", "INVALID_SHIFT", "MISSING_CHECK_IN", "MISSING_CHECK_OUT", "UNLINKED"))
                      else "WARNING" if warnings else "CONFIRMED")
            details.append({
                "teacher_id": record["teacher_id"], "date": day.isoformat(), "weekday": day.weekday(),
                "class_id": shift["class_id"], "center": shift.get("center_name") or shift.get("center_id") or record.get("nearest_center_name"),
                "start_time": str(start) if start else None, "end_time": str(end) if end else None,
                "checked_at": entered.isoformat() if entered else None,
                "check_out_at": exited.isoformat() if exited else None,
                "student_names": check_out.get("student_names") if check_out else None,
                "student_count": len(check_out["student_names"]) if check_out and check_out.get("student_names") is not None else None,
                "result": result, "warnings": warnings, "notes": notes, "minutes": minutes,
                "block_size": len(shifts), "source": inferred or "SAVED",
                "event_pair": (str(check_in["record_id"]) if check_in else None,
                               str(check_out["record_id"]) if check_out else None),
            })

    for record in records:
        if record["event_type"] != "CHECK_IN" or not (record.get("session_block") or str(record["record_id"]) in checkouts):
            continue
        candidates = [out for out in checkouts.get(str(record["record_id"]), [])
                      if out["teacher_id"] == record["teacher_id"]
                      and local_datetime(out["checked_at"]).date() == local_datetime(record["checked_at"]).date()]
        out = max(candidates, key=lambda event: local_datetime(event["checked_at"])) if candidates else None
        used.add(str(record["record_id"]))
        if out:
            used.add(str(out["record_id"]))
        append_session(record, out, bool(out or record.get("session_block")))

    for group in legacy.values():
        remaining = [record for record in group if str(record["record_id"]) not in used]
        ins = [record for record in remaining if record["event_type"] == "CHECK_IN"]
        outs = [record for record in remaining if record["event_type"] == "CHECK_OUT"]
        # Legacy records have no trustworthy block linkage; keep them reviewable, not confirmed.
        if ins or outs:
            entered = max(ins, key=lambda event: local_datetime(event["checked_at"])) if ins else None
            exited = max(outs, key=lambda event: local_datetime(event["checked_at"])) if outs else None
            append_session(entered, exited, False)
            used.update(str(record["record_id"]) for record in remaining)
    for record in records:
        if str(record["record_id"]) not in used:
            append_session(None if record["event_type"] == "CHECK_OUT" else record,
                           record if record["event_type"] == "CHECK_OUT" else None, False)

    # Repeated events must not count the same teacher/date/class twice.
    unique = {}
    rank = {"CONFIRMED": 0, "WARNING": 1, "INVALID": 2, "UNVERIFIED": 3}
    for index, detail in enumerate(details):
        key = (detail["teacher_id"], detail["date"], detail["class_id"] or f"event-{index}")
        if key not in unique or rank[detail["result"]] < rank[unique[key]["result"]]:
            unique[key] = detail
    details = sorted(unique.values(), key=lambda item: (item["teacher_id"], item["date"], item["start_time"] or "", item["class_id"] or ""))
    summary = []
    for teacher in teachers:
        teacher_rows = [item for item in details if item["teacher_id"] == teacher["teacher_id"]]
        confirmed = [item for item in teacher_rows if item["result"] == "CONFIRMED"]
        summary.append({
            **dict(teacher), "confirmed_sessions": len(confirmed),
            "recorded_sessions": sum(bool(item["class_id"]) for item in teacher_rows),
            "warning_sessions": sum(item["result"] == "WARNING" for item in teacher_rows),
            "invalid_sessions": sum(item["result"] == "INVALID" for item in teacher_rows),
            "unverified_events": sum(item["result"] == "UNVERIFIED" for item in teacher_rows),
            "confirmed_minutes": sum(item["minutes"] for item in confirmed),
        })
    # Keep per-class totals, but display a shared check-in/out only once per visit.
    visits = {}
    technical = {"UNKNOWN_SHIFT", "UNLINKED", "LEGACY_SINGLE", "LEGACY_DOUBLE"}
    severity = {"CONFIRMED": 0, "WARNING": 1, "UNVERIFIED": 2, "INVALID": 3}
    for row in details:
        row["warnings"] = [code for code in row["warnings"] if code not in technical]
        row["notes"] = [code for code in row["notes"] if code not in technical]
        key = (row["teacher_id"], row["date"], row["event_pair"])
        if key not in visits:
            visits[key] = {**row, "class_ids": [], "warnings": [], "notes": [], "block_size": 0}
        visit = visits[key]
        if row["class_id"]:
            visit["class_ids"].append(row["class_id"])
            visit["block_size"] += 1
        for field in ("warnings", "notes"):
            visit[field] = list(dict.fromkeys(visit[field] + row[field]))
        if severity[row["result"]] > severity[visit["result"]]:
            visit["result"] = row["result"]
        if row["start_time"]:
            visit["start_time"] = min(visit["start_time"] or row["start_time"], row["start_time"])
        if row["end_time"]:
            visit["end_time"] = max(visit["end_time"] or row["end_time"], row["end_time"])
        visit["class_id"] = ", ".join(visit["class_ids"]) or None
    sessions = sorted(visits.values(), key=lambda row: (row["teacher_id"], row["date"], row["checked_at"] or row["check_out_at"] or ""))
    return {"start_date": start_date.isoformat(), "end_date": end_date.isoformat(), "summary": summary, "details": details, "sessions": sessions}


def build_report_workbook(report, language):
    copy = COPY[language]
    book = Workbook()
    book.remove(book.active)

    def sheet(title, heading, headers):
        base = re.sub(r"[\\/*?:\[\]\x00-\x1f]", "_", title).strip("'") or "Sheet"
        name = base[:31]
        suffix = 2
        while name.casefold() in {value.casefold() for value in book.sheetnames}:
            tail = f" ({suffix})"
            name = base[:31-len(tail)] + tail
            suffix += 1
        result = book.create_sheet(name)
        result.append([heading])
        result.cell(1, 1).data_type = "s"
        result.append([f"{copy['period']}: {report['start_date']} - {report['end_date']}"])
        result.append(headers)
        result.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        result.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        result.cell(1, 1).font = Font(bold=True, size=15, color="842B20")
        for cell in result[3]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="F4DEAA")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        result.freeze_panes = "A4"
        return result

    def finish(result, widths):
        for index, width in enumerate(widths, 1):
            result.column_dimensions[get_column_letter(index)].width = width
        result.auto_filter.ref = f"A3:{get_column_letter(len(widths))}{max(3, result.max_row)}"
        for row in result.iter_rows(min_row=4):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                # Names and notes are user data, never spreadsheet formulas.
                if isinstance(cell.value, str):
                    cell.data_type = "s"

    overview = sheet(copy["summary"], copy["title"], [copy[key] for key in ("code", "teacher", "recorded", "confirmed", "warning", "unknown", "minutes", "invalid")])
    for teacher in report["summary"]:
        overview.append([teacher["teacher_id"], teacher["teacher_name"], teacher["recorded_sessions"], teacher["confirmed_sessions"],
                         teacher["warning_sessions"], teacher["unverified_events"], teacher["confirmed_minutes"], teacher["invalid_sessions"]])
    finish(overview, [18, 30, 22, 22, 22, 24, 28, 22])
    for teacher in report["summary"]:
        rows = [row for row in report["sessions"] if row["teacher_id"] == teacher["teacher_id"]]
        centers = defaultdict(list)
        for row in rows:
            centers[row["center"] or ("Chưa xác định cơ sở" if language == "vi" else "Unknown center")].append(row)
        for center, center_rows in centers.items():
            tab_name = f"{teacher['teacher_name']} - {center}"
            result = sheet(tab_name, f"{teacher['teacher_name']} - {center}",
                           [copy[key] for key in ("date", "weekday", "center", "class", "shift", "in", "out", "block", "result", "notes", "student_count", "student_names")])
            for row in center_rows:
                shift = f"{row['start_time'][:5]} - {row['end_time'][:5]}" if row["start_time"] and row["end_time"] else "-"
                result.append([
                    row["date"], copy["days"][row["weekday"]], row["center"] or "-", row["class_id"] or "-", shift,
                    local_datetime(row["checked_at"]).strftime("%H:%M") if row["checked_at"] else "-",
                    local_datetime(row["check_out_at"]).strftime("%H:%M") if row["check_out_at"] else "-",
                    row["block_size"], copy[row["result"]], "; ".join(copy[code] for code in row["warnings"] + row["notes"]),
                    row["student_count"] if row["student_count"] is not None else copy["students_unknown"],
                    "\n".join(row["student_names"]) if row["student_names"] is not None else copy["students_unknown"],
                ])
                result.cell(result.max_row, 9).fill = PatternFill("solid", fgColor="D9EAD3" if row["result"] == "CONFIRMED" else "FFF2CC")
            finish(result, [15, 17, 24, 28, 22, 14, 14, 20, 22, 65, 20, 45])
            total = sum(row["block_size"] for row in center_rows if row["checked_at"] and row["check_out_at"])
            result.append(["Tổng số ca đã dạy (đủ vào/ra)" if language == "vi" else "Total sessions (checked in/out)", None, None, None, None, None, None, total])
            result.merge_cells(start_row=result.max_row, start_column=1, end_row=result.max_row, end_column=7)
            for cell in result[result.max_row]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="F4DEAA")
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output
