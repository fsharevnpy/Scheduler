from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from math import atan2, cos, radians, sin, sqrt
from uuid import uuid4

from fastapi import APIRouter, Depends, HTTPException, Query
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.orm import Session
from starlette.responses import StreamingResponse

from core.permissions import (
    ROLE_DEVELOPER,
    ROLE_MANAGER,
    ROLE_TEACHER,
    get_current_user,
    require_roles,
)
from db import get_db
from schemas.auth import CurrentUser
from schemas.scheduling import AttendanceCheckInPayload, ManualAttendancePayload
from services.class_assignments import ensure_class_teacher_reassignment_overrides_table, ensure_class_teachers_table


router = APIRouter()
CHECK_IN_COOLDOWN_MINUTES = 5
CHECK_OUT_COOLDOWN_MINUTES = 5
LOCAL_TIMEZONE = "Asia/Ho_Chi_Minh"


EFFECTIVE_ASSIGNMENTS_FOR_CHECKED_AT_CTE = """
            active_teacher_overrides AS (
                SELECT override.class_id, override.original_teacher_id, override.replacement_teacher_id
                FROM class_teacher_reassignment_overrides override
                CROSS JOIN checked
                WHERE override.effective_date <= CAST(checked.local_checked_at AS DATE)
                  AND override.expires_on >= CAST(checked.local_checked_at AS DATE)
            ),
            effective_assignment_rows AS (
                SELECT
                    ct.class_id,
                    COALESCE(active_override.replacement_teacher_id, ct.teacher_id) AS teacher_id
                FROM class_teachers ct
                LEFT JOIN active_teacher_overrides active_override
                    ON active_override.class_id = ct.class_id
                   AND active_override.original_teacher_id = ct.teacher_id

                UNION

                SELECT
                    c.class_id,
                    COALESCE(active_override.replacement_teacher_id, c.teacher_id) AS teacher_id
                FROM classes c
                LEFT JOIN active_teacher_overrides active_override
                    ON active_override.class_id = c.class_id
                   AND active_override.original_teacher_id = c.teacher_id
                WHERE c.teacher_id IS NOT NULL
            ),
            assigned_classes AS (
                SELECT DISTINCT class_id
                FROM effective_assignment_rows
                WHERE teacher_id = :teacher_id
            )
"""


def ensure_attendance_table(db: Session) -> None:
    db.execute(
        text("""
            CREATE TABLE IF NOT EXISTS attendance_records (
                record_id UUID PRIMARY KEY,
                user_code VARCHAR(7) NOT NULL REFERENCES users(user_code) ON DELETE CASCADE,
                teacher_id VARCHAR(20) REFERENCES teachers(teacher_id) ON DELETE SET NULL,
                checked_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                latitude NUMERIC(9,6) NOT NULL,
                longitude NUMERIC(9,6) NOT NULL,
                accuracy_meters NUMERIC(10,2),
                status VARCHAR(10) NOT NULL DEFAULT 'INVALID',
                nearest_center_name VARCHAR(100),
                class_id VARCHAR(20) REFERENCES classes(class_id) ON DELETE SET NULL,
                attendance_status VARCHAR(20) NOT NULL DEFAULT 'UNMATCHED',
                event_type VARCHAR(10) NOT NULL DEFAULT 'CHECK_IN'
            )
        """)
    )
    db.execute(
        text("""
            ALTER TABLE attendance_records
            ADD COLUMN IF NOT EXISTS status VARCHAR(10) NOT NULL DEFAULT 'INVALID'
        """)
    )
    db.execute(
        text("""
            ALTER TABLE attendance_records
            ADD COLUMN IF NOT EXISTS nearest_center_name VARCHAR(100)
        """)
    )
    db.execute(
        text("""
            ALTER TABLE attendance_records
            ADD COLUMN IF NOT EXISTS class_id VARCHAR(20) REFERENCES classes(class_id) ON DELETE SET NULL
        """)
    )
    db.execute(
        text("""
            ALTER TABLE attendance_records
            ADD COLUMN IF NOT EXISTS attendance_status VARCHAR(20) NOT NULL DEFAULT 'UNMATCHED'
        """)
    )
    db.execute(
        text("""
            ALTER TABLE attendance_records
            ADD COLUMN IF NOT EXISTS event_type VARCHAR(10) NOT NULL DEFAULT 'CHECK_IN'
        """)
    )
    db.execute(
        text("""
            UPDATE attendance_records
            SET status = 'INVALID'
            WHERE status = 'NON_VALID'
        """),
    )
    db.execute(
        text("""
            CREATE INDEX IF NOT EXISTS idx_attendance_records_user_time
                ON attendance_records(user_code, checked_at DESC)
        """)
    )
    db.execute(
        text("""
            CREATE INDEX IF NOT EXISTS idx_attendance_records_teacher_time
                ON attendance_records(teacher_id, checked_at DESC)
        """)
    )
    db.commit()


def ensure_assignment_tables(db: Session) -> None:
    ensure_class_teachers_table(db)
    ensure_class_teacher_reassignment_overrides_table(db)


def haversine_distance(lat1, lon1, lat2, lon2):
    r = 6371000.0
    dlat = radians(lat2 - lat1)
    dlon = radians(lon2 - lon1)
    a = sin(dlat / 2) ** 2 + cos(radians(lat1)) * cos(radians(lat2)) * sin(dlon / 2) ** 2
    c = 2 * atan2(sqrt(a), sqrt(1 - a))
    return r * c


def get_closest_center(db: Session, latitude: float, longitude: float):
    centers = db.execute(text("SELECT center_id, name, latitude, longitude FROM centers")).mappings().all()
    closest_distance = None
    closest_center = None

    for center in centers:
        distance = haversine_distance(
            latitude,
            longitude,
            float(center["latitude"]),
            float(center["longitude"]),
        )
        if closest_distance is None or distance < closest_distance:
            closest_distance = distance
            closest_center = center

    if not closest_center:
        return {
            "center_id": None,
            "name": None,
            "distance": None,
            "status": "INVALID",
        }

    return {
        "center_id": closest_center["center_id"],
        "name": closest_center["name"],
        "distance": closest_distance,
        "status": "VALID" if closest_distance is not None and closest_distance <= 200 else "INVALID",
    }


def find_matching_class_session(db: Session, teacher_id: str | None, center_id: str | None, checked_at: datetime):
    if not teacher_id:
        return {"class_id": None, "attendance_status": "UNMATCHED"}

    params = {
        "teacher_id": teacher_id,
        "center_id": center_id,
        "checked_at": checked_at,
    }
    center_filter = "AND c.center_id = :center_id" if center_id else ""

    session = db.execute(
        text(f"""
            WITH checked AS (
                SELECT
                    CAST(:checked_at AS TIMESTAMPTZ) AT TIME ZONE '{LOCAL_TIMEZONE}' AS local_checked_at
            ),
{EFFECTIVE_ASSIGNMENTS_FOR_CHECKED_AT_CTE}
            SELECT
                c.class_id,
                CASE
                    WHEN CAST(checked.local_checked_at AS TIME) < c.start_time - INTERVAL '30 minutes'
                        THEN 'EARLY'
                    WHEN CAST(checked.local_checked_at AS TIME) <= c.start_time
                        THEN 'ON_TIME'
                    WHEN CAST(checked.local_checked_at AS TIME) <= c.end_time
                        THEN 'LATE'
                    ELSE 'UNMATCHED'
                END AS attendance_status
            FROM classes c
            JOIN assigned_classes ac ON ac.class_id = c.class_id
            CROSS JOIN checked
            WHERE c.day_of_week = CASE EXTRACT(ISODOW FROM checked.local_checked_at)
                    WHEN 1 THEN 'T2'
                    WHEN 2 THEN 'T3'
                    WHEN 3 THEN 'T4'
                    WHEN 4 THEN 'T5'
                    WHEN 5 THEN 'T6'
                    WHEN 6 THEN 'T7'
                    ELSE 'CN'
                END
              {center_filter}
              AND CAST(checked.local_checked_at AS TIME) <= c.end_time
            ORDER BY
                CASE
                    WHEN CAST(checked.local_checked_at AS TIME) > c.start_time
                     AND CAST(checked.local_checked_at AS TIME) <= c.end_time THEN 1
                    ELSE 2
                END,
                c.start_time
            LIMIT 1
        """),
        params,
    ).mappings().first()

    if not session:
        return {"class_id": None, "attendance_status": "UNMATCHED"}

    return dict(session)


def get_checkout_status(db: Session, class_id: str, checked_at: datetime) -> str:
    row = db.execute(
        text(f"""
            WITH checked AS (
                SELECT CAST(:checked_at AS TIMESTAMPTZ) AT TIME ZONE '{LOCAL_TIMEZONE}' AS local_checked_at
            )
            SELECT
                CASE
                    WHEN CAST(checked.local_checked_at AS TIME) < c.end_time
                        THEN 'LEFT_EARLY'
                    WHEN CAST(checked.local_checked_at AS TIME) <= c.end_time + INTERVAL '30 minutes'
                        THEN 'COMPLETED'
                    ELSE 'OVERTIME'
                END AS checkout_status
            FROM classes c
            CROSS JOIN checked
            WHERE c.class_id = :class_id
        """),
        {"class_id": class_id, "checked_at": checked_at},
    ).mappings().first()

    return row["checkout_status"] if row else "UNMATCHED"


def find_matching_checkout_session(db: Session, teacher_id: str | None, center_id: str | None, checked_at: datetime):
    if not teacher_id:
        return {"class_id": None, "attendance_status": "UNMATCHED"}

    center_filter = "AND c.center_id = :center_id" if center_id else ""
    session = db.execute(
        text(f"""
            WITH checked AS (
                SELECT CAST(:checked_at AS TIMESTAMPTZ) AT TIME ZONE '{LOCAL_TIMEZONE}' AS local_checked_at
            ),
{EFFECTIVE_ASSIGNMENTS_FOR_CHECKED_AT_CTE}
            SELECT c.class_id
            FROM classes c
            JOIN assigned_classes ac ON ac.class_id = c.class_id
            CROSS JOIN checked
            WHERE c.day_of_week = CASE EXTRACT(ISODOW FROM checked.local_checked_at)
                    WHEN 1 THEN 'T2'
                    WHEN 2 THEN 'T3'
                    WHEN 3 THEN 'T4'
                    WHEN 4 THEN 'T5'
                    WHEN 5 THEN 'T6'
                    WHEN 6 THEN 'T7'
                    ELSE 'CN'
                END
              {center_filter}
              AND CAST(checked.local_checked_at AS TIME) >= c.start_time
            ORDER BY c.start_time DESC
            LIMIT 1
        """),
        {"teacher_id": teacher_id, "center_id": center_id, "checked_at": checked_at},
    ).mappings().first()

    if not session:
        return {"class_id": None, "attendance_status": "UNMATCHED"}

    return {
        "class_id": session["class_id"],
        "attendance_status": get_checkout_status(db, session["class_id"], checked_at),
    }


def find_checkout_class_for_open_check_in(db: Session, teacher_id: str | None, check_in_class_id: str, checked_at: datetime):
    if not teacher_id:
        return check_in_class_id

    row = db.execute(
        text(f"""
            WITH RECURSIVE checked AS (
                SELECT CAST(:checked_at AS TIMESTAMPTZ) AT TIME ZONE '{LOCAL_TIMEZONE}' AS local_checked_at
            ),
            active_teacher_overrides AS (
                SELECT override.class_id, override.original_teacher_id, override.replacement_teacher_id
                FROM class_teacher_reassignment_overrides override
                CROSS JOIN checked
                WHERE override.effective_date <= CAST(checked.local_checked_at AS DATE)
                  AND override.expires_on >= CAST(checked.local_checked_at AS DATE)
            ),
            effective_assignment_rows AS (
                SELECT
                    ct.class_id,
                    COALESCE(active_override.replacement_teacher_id, ct.teacher_id) AS teacher_id
                FROM class_teachers ct
                LEFT JOIN active_teacher_overrides active_override
                    ON active_override.class_id = ct.class_id
                   AND active_override.original_teacher_id = ct.teacher_id

                UNION

                SELECT
                    c.class_id,
                    COALESCE(active_override.replacement_teacher_id, c.teacher_id) AS teacher_id
                FROM classes c
                LEFT JOIN active_teacher_overrides active_override
                    ON active_override.class_id = c.class_id
                   AND active_override.original_teacher_id = c.teacher_id
                WHERE c.teacher_id IS NOT NULL
            ),
            assigned_classes AS (
                SELECT DISTINCT class_id
                FROM effective_assignment_rows
                WHERE teacher_id = :teacher_id
            ),
            contiguous_classes AS (
                SELECT c.class_id, c.center_id, c.day_of_week, c.start_time, c.end_time
                FROM classes c
                WHERE c.class_id = :check_in_class_id

                UNION ALL

                SELECT next_class.class_id,
                       next_class.center_id,
                       next_class.day_of_week,
                       next_class.start_time,
                       next_class.end_time
                FROM classes next_class
                JOIN contiguous_classes current_class
                  ON next_class.center_id = current_class.center_id
                 AND next_class.day_of_week = current_class.day_of_week
                 AND next_class.start_time = current_class.end_time
                JOIN assigned_classes ac ON ac.class_id = next_class.class_id
            )
            SELECT cc.class_id
            FROM contiguous_classes cc
            CROSS JOIN checked
            ORDER BY
                CASE
                    WHEN CAST(checked.local_checked_at AS TIME) >= cc.start_time THEN 0
                    ELSE 1
                END,
                CASE
                    WHEN CAST(checked.local_checked_at AS TIME) >= cc.start_time THEN cc.start_time
                    ELSE NULL
                END DESC NULLS LAST,
                cc.start_time
            LIMIT 1
        """),
        {
            "teacher_id": teacher_id,
            "check_in_class_id": check_in_class_id,
            "checked_at": checked_at,
        },
    ).mappings().first()

    return row["class_id"] if row else check_in_class_id


def find_open_check_in(db: Session, user_code: str):
    row = db.execute(
        text(f"""
            WITH checked AS (
                SELECT NOW() AT TIME ZONE '{LOCAL_TIMEZONE}' AS local_checked_at
            )
            SELECT ar.class_id
            FROM attendance_records ar
            JOIN classes c ON c.class_id = ar.class_id
            CROSS JOIN checked
            WHERE ar.user_code = :user_code
              AND ar.event_type = 'CHECK_IN'
              AND ar.class_id IS NOT NULL
              AND (ar.checked_at AT TIME ZONE '{LOCAL_TIMEZONE}')::date = checked.local_checked_at::date
              AND NOT EXISTS (
                  SELECT 1
                  FROM attendance_records checkout
                  WHERE checkout.user_code = ar.user_code
                    AND checkout.event_type = 'CHECK_OUT'
                    AND checkout.checked_at > ar.checked_at
                    AND (checkout.checked_at AT TIME ZONE '{LOCAL_TIMEZONE}')::date = checked.local_checked_at::date
              )
            ORDER BY ar.checked_at DESC
            LIMIT 1
        """),
        {"user_code": user_code},
    ).mappings().first()

    return row["class_id"] if row else None


def format_report_date(value: date) -> str:
    return value.strftime("%d/%m/%Y")


def format_report_time(value) -> str:
    if not value:
        return ""

    return value.strftime("%H:%M") if hasattr(value, "strftime") else str(value)[:5]


def get_week_ranges(start_date: date, end_date: date):
    ranges = []
    current_start = start_date

    while current_start <= end_date:
        week_end = min(end_date, current_start + timedelta(days=6 - current_start.weekday()))
        ranges.append((current_start, week_end))
        current_start = week_end + timedelta(days=1)

    return ranges


def is_valid_check_in(record) -> bool:
    return bool(record) and record["status"] in ("VALID", "MANUAL") and record["attendance_status"] in ("ON_TIME", "EARLY", "MANUAL")


def is_valid_check_out(record) -> bool:
    return bool(record) and record["status"] in ("VALID", "MANUAL") and record["attendance_status"] in ("COMPLETED", "MANUAL")


def get_attendance_marker(valid_record, any_record) -> str:
    if valid_record:
        return "OK"

    return "ISSUE" if any_record else "MISSING"


def get_overall_attendance_marker(check_in_marker: str, check_out_marker: str) -> str:
    if check_in_marker == "OK" and check_out_marker == "OK":
        return "OK"

    if check_in_marker == "MISSING" or check_out_marker == "MISSING":
        return "MISSING"

    return "ISSUE"


def enrich_weekly_schedule_rows(weekly_schedule_rows, attendance_rows):
    enriched_rows = [dict(row) for row in weekly_schedule_rows]
    attendance_by_teacher_date_class = {}

    for record in attendance_rows:
        key = (record["teacher_id"], record["attendance_date"], record["class_id"])
        if key not in attendance_by_teacher_date_class:
            attendance_by_teacher_date_class[key] = {
                "check_ins": [],
                "check_outs": [],
            }

        target = attendance_by_teacher_date_class[key]["check_outs" if record["event_type"] == "CHECK_OUT" else "check_ins"]
        target.append(record)

    schedule_groups = {}
    for row in enriched_rows:
        key = (row["teacher_id"], row["teaching_date"], row["center_id"])
        schedule_groups.setdefault(key, []).append(row)

    for group_key, group_rows in schedule_groups.items():
        group_rows.sort(key=lambda item: (item["start_time"], item["end_time"], item["class_id"]))
        current_block = []
        blocks = []

        for row in group_rows:
            if current_block and format_report_time(current_block[-1]["end_time"]) != format_report_time(row["start_time"]):
                blocks.append(current_block)
                current_block = []
            current_block.append(row)

        if current_block:
            blocks.append(current_block)

        teacher_id, teaching_date, _center_id = group_key
        for block in blocks:
            class_ids = [row["class_id"] for row in block]
            block_check_ins = []
            block_check_outs = []

            for class_id in class_ids:
                records = attendance_by_teacher_date_class.get((teacher_id, teaching_date, class_id), {})
                block_check_ins.extend(records.get("check_ins", []))
                block_check_outs.extend(records.get("check_outs", []))

            valid_check_in = next((record for record in block_check_ins if is_valid_check_in(record)), None)
            valid_check_out = next((record for record in block_check_outs if is_valid_check_out(record)), None)
            check_in_record = valid_check_in or (block_check_ins[0] if block_check_ins else None)
            check_out_record = valid_check_out or (block_check_outs[0] if block_check_outs else None)
            check_in_marker = get_attendance_marker(valid_check_in, check_in_record)
            check_out_marker = get_attendance_marker(valid_check_out, check_out_record)
            overall_marker = get_overall_attendance_marker(check_in_marker, check_out_marker)

            for row in block:
                row["check_in_marker"] = check_in_marker
                row["check_in_detail"] = (
                    f"{check_in_marker} - {format_report_time(check_in_record['checked_time'])} ({check_in_record['attendance_status']})"
                    if check_in_record
                    else "MISSING"
                )
                row["check_out_marker"] = check_out_marker
                row["check_out_detail"] = (
                    f"{check_out_marker} - {format_report_time(check_out_record['checked_time'])} ({check_out_record['attendance_status']})"
                    if check_out_record
                    else "MISSING"
                )
                row["overall_marker"] = overall_marker

    return enriched_rows


def build_teacher_attendance_workbook(rows, weekly_schedule_rows, attendance_rows, start_date: date, end_date: date) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Teacher Attendance"

    headers = [
        "Mã giáo viên",
        "Tên giáo viên",
        "Số ca dạy",
        "Check-in đúng giờ",
        "Check-in sớm",
        "Check-in trễ",
        "Check-in hợp lệ",
        "Check-in không hợp lệ",
        "Tổng check-in",
        "Checkout hoàn thành",
        "Về sớm",
        "Overtime",
        "Check-out hợp lệ",
        "Check-out không hợp lệ",
        "Tổng checkout",
        "Manual",
        "Tổng sự kiện",
    ]

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    sheet.cell(row=1, column=1).value = f"Khoảng thời gian: {format_report_date(start_date)} - {format_report_date(end_date)}"
    sheet.cell(row=1, column=1).font = Font(bold=True, size=13)
    sheet.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sheet.cell(row=2, column=1).value = "Báo cáo điểm danh giáo viên - cơ sở tính lương"
    sheet.cell(row=2, column=1).font = Font(bold=True)
    sheet.cell(row=2, column=1).alignment = Alignment(horizontal="center")

    header_fill = PatternFill("solid", fgColor="D9EAD3")
    valid_count_fill = PatternFill("solid", fgColor="B6D7A8")
    valid_count_columns = {7, 13}
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=4, column=column_index)
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = valid_count_fill if column_index in valid_count_columns else header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index, row in enumerate(rows, start=5):
        values = [
            row["user_code"] or row["teacher_id"],
            row["teacher_name"],
            row["scheduled_sessions"],
            row["on_time_check_in"],
            row["early_check_in"],
            row["late_check_in"],
            row["valid_check_in"],
            row["invalid_check_in"],
            row["total_check_in"],
            row["completed_check_out"],
            row["left_early_check_out"],
            row["overtime_check_out"],
            row["valid_check_out"],
            row["invalid_check_out"],
            row["total_check_out"],
            row["manual_events"],
            row["total_events"],
        ]

        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.value = value
            cell.alignment = Alignment(horizontal="left" if column_index <= 2 else "center")
            if column_index in valid_count_columns:
                cell.fill = valid_count_fill
                cell.font = Font(bold=True)

    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{max(4, len(rows) + 4)}"

    widths = [14, 28, 12, 16, 14, 14, 16, 18, 14, 19, 12, 12, 17, 19, 14, 12, 14]
    for column_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    weekly_schedule_rows = enrich_weekly_schedule_rows(weekly_schedule_rows, attendance_rows)
    schedule_by_week = {}
    for week_index, (week_start, week_end) in enumerate(get_week_ranges(start_date, end_date), start=1):
        schedule_by_week[week_index] = {
            "start": week_start,
            "end": week_end,
            "rows": [],
        }

    for schedule_row in weekly_schedule_rows:
        teaching_date = schedule_row["teaching_date"]
        for week_index, week in schedule_by_week.items():
            if week["start"] <= teaching_date <= week["end"]:
                week["rows"].append(schedule_row)
                break

    weekly_headers = [
        "Ngày",
        "Thứ",
        "Mã giáo viên",
        "Tên giáo viên",
        "Cơ sở",
        "Bắt đầu",
        "Kết thúc",
        "Lớp/Ca",
        "Check-in",
        "Checkout",
        "Kết quả",
    ]
    weekly_widths = [14, 10, 14, 28, 18, 12, 12, 22, 28, 28, 14]
    marker_fills = {
        "OK": PatternFill("solid", fgColor="D9EAD3"),
        "MISSING": PatternFill("solid", fgColor="FFF2CC"),
        "ISSUE": PatternFill("solid", fgColor="F4CCCC"),
    }

    for week_index, week in schedule_by_week.items():
        week_sheet = workbook.create_sheet(
            title=f"Week {week_index} {week['start'].strftime('%d-%m')}"
        )
        week_sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(weekly_headers))
        week_sheet.cell(row=1, column=1).value = (
            f"Tuần {week_index}: {format_report_date(week['start'])} - {format_report_date(week['end'])}"
        )
        week_sheet.cell(row=1, column=1).font = Font(bold=True, size=13)
        week_sheet.cell(row=1, column=1).alignment = Alignment(horizontal="center")

        legend_items = [
            ("OK", "Có check-in và checkout hợp lệ"),
            ("MISSING", "Thiếu check-in hoặc checkout"),
            ("ISSUE", "Có dữ liệu nhưng bị trễ/về sớm/overtime/invalid"),
        ]
        for legend_index, (marker, description) in enumerate(legend_items, start=1):
            marker_cell = week_sheet.cell(row=2, column=(legend_index - 1) * 3 + 1)
            marker_cell.value = marker
            marker_cell.font = Font(bold=True)
            marker_cell.fill = marker_fills[marker]
            marker_cell.alignment = Alignment(horizontal="center")

            description_cell = week_sheet.cell(row=2, column=(legend_index - 1) * 3 + 2)
            description_cell.value = description
            description_cell.alignment = Alignment(horizontal="left")

        for column_index, header in enumerate(weekly_headers, start=1):
            cell = week_sheet.cell(row=3, column=column_index)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        if week["rows"]:
            for row_index, schedule_row in enumerate(week["rows"], start=4):
                values = [
                    format_report_date(schedule_row["teaching_date"]),
                    schedule_row["day_of_week"],
                    schedule_row["user_code"] or schedule_row["teacher_id"],
                    schedule_row["teacher_name"],
                    schedule_row["center_name"],
                    format_report_time(schedule_row["start_time"]),
                    format_report_time(schedule_row["end_time"]),
                    schedule_row["class_id"],
                    schedule_row["check_in_detail"],
                    schedule_row["check_out_detail"],
                    schedule_row["overall_marker"],
                ]
                for column_index, value in enumerate(values, start=1):
                    cell = week_sheet.cell(row=row_index, column=column_index)
                    cell.value = value
                    cell.alignment = Alignment(horizontal="left" if column_index in (3, 4, 5, 8, 9, 10) else "center")
                    if column_index == 11:
                        cell.fill = marker_fills.get(schedule_row["overall_marker"], PatternFill())
                        cell.font = Font(bold=True)
        else:
            week_sheet.cell(row=4, column=1).value = "Không có ca dạy trong tuần này"
            week_sheet.cell(row=4, column=1).alignment = Alignment(horizontal="center")
            week_sheet.merge_cells(start_row=4, start_column=1, end_row=4, end_column=len(weekly_headers))

        week_sheet.freeze_panes = "A4"
        week_sheet.auto_filter.ref = f"A3:{get_column_letter(len(weekly_headers))}{max(3, len(week['rows']) + 3)}"
        for column_index, width in enumerate(weekly_widths, start=1):
            week_sheet.column_dimensions[get_column_letter(column_index)].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


@router.get("/attendance-records")
def get_attendance_records(
    current_user: CurrentUser = Depends(require_roles(ROLE_DEVELOPER, ROLE_MANAGER, ROLE_TEACHER)),
    db: Session = Depends(get_db),
):
    ensure_attendance_table(db)

    params = {}
    query = """
        SELECT
            ar.record_id,
            ar.user_code,
            u.display_name,
            ar.teacher_id,
            ar.checked_at,
            ar.latitude,
            ar.longitude,
            ar.accuracy_meters,
            ar.status,
            ar.nearest_center_name,
            ar.class_id,
            ar.attendance_status,
            ar.event_type
        FROM attendance_records ar
        JOIN users u ON u.user_code = ar.user_code
    """

    if current_user.role == ROLE_TEACHER:
        query += " WHERE ar.user_code = :user_code"
        params["user_code"] = current_user.user_code

    query += " ORDER BY ar.checked_at DESC LIMIT 200"

    records = db.execute(text(query), params).mappings().all()
    return [dict(item) for item in records]


@router.get("/attendance-records/export")
def export_teacher_attendance_report(
    start_date: date = Query(...),
    end_date: date = Query(...),
    _: CurrentUser = Depends(require_roles(ROLE_DEVELOPER, ROLE_MANAGER)),
    db: Session = Depends(get_db),
):
    ensure_attendance_table(db)
    ensure_assignment_tables(db)

    if end_date < start_date:
        raise HTTPException(status_code=400, detail="end_date must be after or equal to start_date")

    if (end_date - start_date).days > 366:
        raise HTTPException(status_code=400, detail="date range cannot be longer than 366 days")

    rows = db.execute(
        text(f"""
            WITH report_records AS (
                SELECT *
                FROM attendance_records ar
                WHERE (ar.checked_at AT TIME ZONE '{LOCAL_TIMEZONE}')::date BETWEEN :start_date AND :end_date
            ),
            report_dates AS (
                SELECT
                    day::date AS report_date,
                    CASE EXTRACT(ISODOW FROM day)
                        WHEN 1 THEN 'T2'
                        WHEN 2 THEN 'T3'
                        WHEN 3 THEN 'T4'
                        WHEN 4 THEN 'T5'
                        WHEN 5 THEN 'T6'
                        WHEN 6 THEN 'T7'
                        ELSE 'CN'
                    END AS day_of_week
                FROM generate_series(CAST(:start_date AS date), CAST(:end_date AS date), INTERVAL '1 day') AS day
            ),
            teacher_classes AS (
                SELECT
                    rd.report_date,
                    c.class_id,
                    c.day_of_week,
                    c.center_id,
                    c.start_time,
                    c.end_time,
                    COALESCE(active_override.replacement_teacher_id, ct.teacher_id) AS teacher_id
                FROM report_dates rd
                JOIN classes c ON c.day_of_week = rd.day_of_week
                JOIN class_teachers ct ON ct.class_id = c.class_id
                LEFT JOIN class_teacher_reassignment_overrides active_override
                    ON active_override.class_id = c.class_id
                   AND active_override.original_teacher_id = ct.teacher_id
                   AND active_override.effective_date <= rd.report_date
                   AND active_override.expires_on >= rd.report_date
                WHERE c.status <> 'CANCELLED'

                UNION

                SELECT
                    rd.report_date,
                    c.class_id,
                    c.day_of_week,
                    c.center_id,
                    c.start_time,
                    c.end_time,
                    COALESCE(active_override.replacement_teacher_id, c.teacher_id) AS teacher_id
                FROM report_dates rd
                JOIN classes c ON c.day_of_week = rd.day_of_week
                LEFT JOIN class_teacher_reassignment_overrides active_override
                    ON active_override.class_id = c.class_id
                   AND active_override.original_teacher_id = c.teacher_id
                   AND active_override.effective_date <= rd.report_date
                   AND active_override.expires_on >= rd.report_date
                WHERE c.teacher_id IS NOT NULL
                  AND c.status <> 'CANCELLED'
            ),
            scheduled_sessions AS (
                SELECT teacher_id, COUNT(*) AS scheduled_sessions
                FROM teacher_classes
                GROUP BY teacher_id
            )
            SELECT
                t.teacher_id,
                t.name AS teacher_name,
                u.user_code,
                COALESCE(MAX(ss.scheduled_sessions), 0) AS scheduled_sessions,
                COUNT(rr.record_id) AS total_events,
                COUNT(rr.record_id) FILTER (WHERE rr.event_type = 'CHECK_IN') AS total_check_in,
                COUNT(rr.record_id) FILTER (
                    WHERE rr.event_type = 'CHECK_IN'
                      AND rr.status IN ('VALID', 'MANUAL')
                      AND rr.attendance_status IN ('ON_TIME', 'EARLY', 'MANUAL')
                ) AS valid_check_in,
                COUNT(rr.record_id) FILTER (
                    WHERE rr.event_type = 'CHECK_IN'
                      AND (
                        rr.status = 'INVALID'
                        OR rr.attendance_status IN ('UNMATCHED', 'MISSED')
                      )
                ) AS invalid_check_in,
                COUNT(rr.record_id) FILTER (
                    WHERE rr.event_type = 'CHECK_IN'
                      AND rr.attendance_status = 'ON_TIME'
                ) AS on_time_check_in,
                COUNT(rr.record_id) FILTER (
                    WHERE rr.event_type = 'CHECK_IN'
                      AND rr.attendance_status = 'EARLY'
                ) AS early_check_in,
                COUNT(rr.record_id) FILTER (
                    WHERE rr.event_type = 'CHECK_IN'
                      AND rr.attendance_status = 'LATE'
                ) AS late_check_in,
                COUNT(rr.record_id) FILTER (WHERE rr.event_type = 'CHECK_OUT') AS total_check_out,
                COUNT(rr.record_id) FILTER (
                    WHERE rr.event_type = 'CHECK_OUT'
                      AND rr.attendance_status = 'COMPLETED'
                ) AS completed_check_out,
                COUNT(rr.record_id) FILTER (
                    WHERE rr.event_type = 'CHECK_OUT'
                      AND rr.status IN ('VALID', 'MANUAL')
                      AND rr.attendance_status IN ('COMPLETED', 'MANUAL')
                ) AS valid_check_out,
                COUNT(rr.record_id) FILTER (
                    WHERE rr.event_type = 'CHECK_OUT'
                      AND (
                        rr.status = 'INVALID'
                        OR rr.attendance_status = 'UNMATCHED'
                      )
                ) AS invalid_check_out,
                COUNT(rr.record_id) FILTER (
                    WHERE rr.event_type = 'CHECK_OUT'
                      AND rr.attendance_status = 'LEFT_EARLY'
                ) AS left_early_check_out,
                COUNT(rr.record_id) FILTER (
                    WHERE rr.event_type = 'CHECK_OUT'
                      AND rr.attendance_status = 'OVERTIME'
                ) AS overtime_check_out,
                COUNT(rr.record_id) FILTER (
                    WHERE rr.status = 'MANUAL'
                       OR rr.attendance_status = 'MANUAL'
                ) AS manual_events
            FROM teachers t
            LEFT JOIN users u ON u.teacher_id = t.teacher_id AND u.role = 'teacher'
            LEFT JOIN report_records rr ON rr.teacher_id = t.teacher_id
            LEFT JOIN scheduled_sessions ss ON ss.teacher_id = t.teacher_id
            GROUP BY t.teacher_id, t.name, u.user_code
            ORDER BY COALESCE(u.user_code, t.teacher_id), t.name
        """),
        {
            "start_date": start_date,
            "end_date": end_date,
        },
    ).mappings().all()

    weekly_schedule_rows = db.execute(
        text("""
            WITH report_dates AS (
                SELECT
                    day::date AS teaching_date,
                    CASE EXTRACT(ISODOW FROM day)
                        WHEN 1 THEN 'T2'
                        WHEN 2 THEN 'T3'
                        WHEN 3 THEN 'T4'
                        WHEN 4 THEN 'T5'
                        WHEN 5 THEN 'T6'
                        WHEN 6 THEN 'T7'
                        ELSE 'CN'
                    END AS day_of_week
                FROM generate_series(CAST(:start_date AS date), CAST(:end_date AS date), INTERVAL '1 day') AS day
            ),
            teacher_classes AS (
                SELECT
                    rd.teaching_date,
                    c.class_id,
                    c.day_of_week,
                    c.center_id,
                    c.start_time,
                    c.end_time,
                    COALESCE(active_override.replacement_teacher_id, ct.teacher_id) AS teacher_id
                FROM report_dates rd
                JOIN classes c ON c.day_of_week = rd.day_of_week
                JOIN class_teachers ct ON ct.class_id = c.class_id
                LEFT JOIN class_teacher_reassignment_overrides active_override
                    ON active_override.class_id = c.class_id
                   AND active_override.original_teacher_id = ct.teacher_id
                   AND active_override.effective_date <= rd.teaching_date
                   AND active_override.expires_on >= rd.teaching_date
                WHERE c.status <> 'CANCELLED'

                UNION

                SELECT
                    rd.teaching_date,
                    c.class_id,
                    c.day_of_week,
                    c.center_id,
                    c.start_time,
                    c.end_time,
                    COALESCE(active_override.replacement_teacher_id, c.teacher_id) AS teacher_id
                FROM report_dates rd
                JOIN classes c ON c.day_of_week = rd.day_of_week
                LEFT JOIN class_teacher_reassignment_overrides active_override
                    ON active_override.class_id = c.class_id
                   AND active_override.original_teacher_id = c.teacher_id
                   AND active_override.effective_date <= rd.teaching_date
                   AND active_override.expires_on >= rd.teaching_date
                WHERE c.teacher_id IS NOT NULL
                  AND c.status <> 'CANCELLED'
            )
            SELECT
                rd.teaching_date,
                rd.day_of_week,
                tc.class_id,
                tc.start_time,
                tc.end_time,
                tc.teacher_id,
                t.name AS teacher_name,
                u.user_code,
                tc.center_id,
                centers.name AS center_name
            FROM report_dates rd
            JOIN teacher_classes tc ON tc.teaching_date = rd.teaching_date
            JOIN teachers t ON t.teacher_id = tc.teacher_id
            LEFT JOIN users u ON u.teacher_id = t.teacher_id AND u.role = 'teacher'
            JOIN centers ON centers.center_id = tc.center_id
            ORDER BY
                rd.teaching_date,
                tc.start_time,
                centers.center_id,
                COALESCE(u.user_code, tc.teacher_id),
                t.name
        """),
        {
            "start_date": start_date,
            "end_date": end_date,
        },
    ).mappings().all()

    attendance_rows = db.execute(
        text(f"""
            SELECT
                ar.teacher_id,
                (ar.checked_at AT TIME ZONE '{LOCAL_TIMEZONE}')::date AS attendance_date,
                CAST(ar.checked_at AT TIME ZONE '{LOCAL_TIMEZONE}' AS TIME) AS checked_time,
                ar.class_id,
                ar.event_type,
                ar.status,
                ar.attendance_status
            FROM attendance_records ar
            WHERE ar.teacher_id IS NOT NULL
              AND ar.class_id IS NOT NULL
              AND (ar.checked_at AT TIME ZONE '{LOCAL_TIMEZONE}')::date BETWEEN :start_date AND :end_date
            ORDER BY ar.checked_at
        """),
        {
            "start_date": start_date,
            "end_date": end_date,
        },
    ).mappings().all()

    report_file = build_teacher_attendance_workbook(rows, weekly_schedule_rows, attendance_rows, start_date, end_date)
    filename = f"teacher_attendance_{start_date.isoformat()}_to_{end_date.isoformat()}.xlsx"
    return StreamingResponse(
        report_file,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/attendance-records/check-in")
def check_in_attendance(
    payload: AttendanceCheckInPayload,
    current_user: CurrentUser = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    ensure_attendance_table(db)
    ensure_assignment_tables(db)

    recent_check_in = db.execute(
        text("""
            SELECT checked_at
            FROM attendance_records
            WHERE user_code = :user_code
              AND event_type = 'CHECK_IN'
              AND checked_at > NOW() - (:cooldown_minutes * INTERVAL '1 minute')
            ORDER BY checked_at DESC
            LIMIT 1
        """),
        {
            "user_code": current_user.user_code,
            "cooldown_minutes": CHECK_IN_COOLDOWN_MINUTES,
        },
    ).mappings().first()

    if recent_check_in:
        raise HTTPException(
            status_code=429,
            detail=f"please wait {CHECK_IN_COOLDOWN_MINUTES} minutes between check-ins",
        )

    record_id = str(uuid4())
    record_checked_at = datetime.now(timezone.utc)
    closest_center = get_closest_center(db, payload.latitude, payload.longitude)
    session_match = find_matching_class_session(db, current_user.teacher_id, closest_center["center_id"], record_checked_at)

    record = db.execute(
        text("""
            INSERT INTO attendance_records (
                record_id, user_code, teacher_id, checked_at, latitude, longitude, accuracy_meters,
                status, nearest_center_name, class_id, attendance_status, event_type
            )
            VALUES (
                :record_id, :user_code, :teacher_id, :checked_at, :latitude, :longitude, :accuracy_meters,
                :status, :nearest_center_name, :class_id, :attendance_status, 'CHECK_IN'
            )
            RETURNING
                record_id,
                user_code,
                teacher_id,
                checked_at,
                latitude,
                longitude,
                accuracy_meters,
                status,
                nearest_center_name,
                class_id,
                attendance_status,
                event_type
        """),
        {
            "record_id": record_id,
            "user_code": current_user.user_code,
            "teacher_id": current_user.teacher_id,
            "checked_at": record_checked_at,
            "latitude": payload.latitude,
            "longitude": payload.longitude,
            "accuracy_meters": payload.accuracy_meters,
            "status": closest_center["status"],
            "nearest_center_name": closest_center["name"],
            "class_id": session_match["class_id"],
            "attendance_status": session_match["attendance_status"],
        },
    ).mappings().first()
    db.commit()

    return {
        **dict(record),
        "display_name": current_user.display_name,
    }


@router.post("/attendance-records/check-out")
def check_out_attendance(
    payload: AttendanceCheckInPayload,
    current_user: CurrentUser = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    ensure_attendance_table(db)
    ensure_assignment_tables(db)

    recent_check_out = db.execute(
        text("""
            SELECT checked_at
            FROM attendance_records
            WHERE user_code = :user_code
              AND event_type = 'CHECK_OUT'
              AND checked_at > NOW() - (:cooldown_minutes * INTERVAL '1 minute')
            ORDER BY checked_at DESC
            LIMIT 1
        """),
        {
            "user_code": current_user.user_code,
            "cooldown_minutes": CHECK_OUT_COOLDOWN_MINUTES,
        },
    ).mappings().first()

    if recent_check_out:
        raise HTTPException(
            status_code=429,
            detail=f"please wait {CHECK_OUT_COOLDOWN_MINUTES} minutes between check-outs",
        )

    check_in_class_id = find_open_check_in(db, current_user.user_code)
    if not check_in_class_id:
        raise HTTPException(status_code=400, detail="no open check-in found for today")

    record_id = str(uuid4())
    record_checked_at = datetime.now(timezone.utc)
    closest_center = get_closest_center(db, payload.latitude, payload.longitude)
    class_id = find_checkout_class_for_open_check_in(
        db,
        current_user.teacher_id,
        check_in_class_id,
        record_checked_at,
    )
    checkout_status = get_checkout_status(db, class_id, record_checked_at)

    record = db.execute(
        text("""
            INSERT INTO attendance_records (
                record_id, user_code, teacher_id, checked_at, latitude, longitude, accuracy_meters,
                status, nearest_center_name, class_id, attendance_status, event_type
            )
            VALUES (
                :record_id, :user_code, :teacher_id, :checked_at, :latitude, :longitude, :accuracy_meters,
                :status, :nearest_center_name, :class_id, :attendance_status, 'CHECK_OUT'
            )
            RETURNING
                record_id,
                user_code,
                teacher_id,
                checked_at,
                latitude,
                longitude,
                accuracy_meters,
                status,
                nearest_center_name,
                class_id,
                attendance_status,
                event_type
        """),
        {
            "record_id": record_id,
            "user_code": current_user.user_code,
            "teacher_id": current_user.teacher_id,
            "checked_at": record_checked_at,
            "latitude": payload.latitude,
            "longitude": payload.longitude,
            "accuracy_meters": payload.accuracy_meters,
            "status": closest_center["status"],
            "nearest_center_name": closest_center["name"],
            "class_id": class_id,
            "attendance_status": checkout_status,
        },
    ).mappings().first()
    db.commit()

    return {
        **dict(record),
        "display_name": current_user.display_name,
    }


@router.post("/attendance-records/manual")
def create_manual_attendance(
    payload: ManualAttendancePayload,
    current_user: CurrentUser = Depends(require_roles(ROLE_DEVELOPER, ROLE_MANAGER)),
    db: Session = Depends(get_db),
):
    ensure_attendance_table(db)
    ensure_assignment_tables(db)

    teacher = db.execute(
        text("""
            SELECT t.teacher_id, t.name, u.user_code, u.display_name
            FROM teachers t
            JOIN users u ON u.teacher_id = t.teacher_id AND u.role = 'teacher'
            WHERE t.teacher_id = :teacher_id
              AND t.active = TRUE
              AND u.active = TRUE
            LIMIT 1
        """),
        {"teacher_id": payload.teacher_id},
    ).mappings().first()

    if not teacher:
        raise HTTPException(status_code=404, detail="teacher not found")

    center = db.execute(
        text("""
            SELECT center_id, name, latitude, longitude
            FROM centers
            WHERE center_id = :center_id
            LIMIT 1
        """),
        {"center_id": payload.center_id},
    ).mappings().first()

    if not center:
        raise HTTPException(status_code=404, detail="center not found")

    record_id = str(uuid4())
    session_match = (
        find_matching_checkout_session(db, teacher["teacher_id"], center["center_id"], payload.checked_at)
        if payload.event_type == "CHECK_OUT"
        else find_matching_class_session(db, teacher["teacher_id"], center["center_id"], payload.checked_at)
    )
    record = db.execute(
        text("""
            INSERT INTO attendance_records (
                record_id,
                user_code,
                teacher_id,
                checked_at,
                latitude,
                longitude,
                accuracy_meters,
                status,
                nearest_center_name,
                class_id,
                attendance_status,
                event_type
            )
            VALUES (
                :record_id,
                :user_code,
                :teacher_id,
                :checked_at,
                :latitude,
                :longitude,
                NULL,
                :record_status,
                :nearest_center_name,
                :class_id,
                :attendance_status,
                :event_type
            )
            RETURNING
                record_id,
                user_code,
                teacher_id,
                checked_at,
                latitude,
                longitude,
                accuracy_meters,
                status,
                nearest_center_name,
                class_id,
                attendance_status,
                event_type
        """),
        {
            "record_id": record_id,
            "user_code": teacher["user_code"],
            "teacher_id": teacher["teacher_id"],
            "checked_at": payload.checked_at,
            "latitude": center["latitude"],
            "longitude": center["longitude"],
            "record_status": "MANUAL",
            "nearest_center_name": center["name"],
            "class_id": session_match["class_id"],
            "attendance_status": "MANUAL" if payload.event_type == "CHECK_IN" else session_match["attendance_status"],
            "event_type": payload.event_type,
        },
    ).mappings().first()
    db.commit()

    return {
        **dict(record),
        "display_name": teacher["display_name"],
    }
