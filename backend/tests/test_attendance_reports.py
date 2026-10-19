import unittest
import json
from datetime import date, datetime, timezone, timedelta
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook
from fastapi import HTTPException
from services.attendance_reports import build_attendance_report, build_report_workbook
from api.routes.attendance import resolve_checkout_session, snapshot_checkout, snapshot_attendance_session, load_attendance_report, find_open_check_in

DAY = date(2026, 10, 19)
TEACHERS = [{"teacher_id": "T1", "teacher_name": "Teacher One", "user_code": "TC00001"}]
BLOCK = [
    {"class_id": "C1", "center_id": "CS1", "center_name": "Center", "start_time": "07:00:00", "end_time": "08:30:00"},
    {"class_id": "C2", "center_id": "CS1", "center_name": "Center", "start_time": "08:30:00", "end_time": "11:00:00"},
]


def pair(entered="07:00:00", exited="11:00:00", **extra):
    base = {"teacher_id": "T1", "user_code": "TC00001", "status": "VALID"}
    return [
        {**base, "record_id": "in1", "event_type": "CHECK_IN", "class_id": "C1",
         "checked_at": f"2026-10-19T{entered}+07:00", "attendance_status": "ON_TIME", "session_block": BLOCK, **extra},
        {**base, "record_id": "out1", "event_type": "CHECK_OUT", "class_id": "C2",
         "checked_at": f"2026-10-19T{exited}+07:00", "attendance_status": "COMPLETED", "check_in_record_id": "in1"},
    ]


class AttendanceReportTests(unittest.TestCase):
    def report(self, records):
        return build_attendance_report(TEACHERS, records, DAY, DAY)

    def test_double_shift_counts_two_classes_not_one_event_pair(self):
        report = self.report(pair())
        self.assertEqual(report["summary"][0]["confirmed_sessions"], 2)
        self.assertEqual(report["summary"][0]["confirmed_minutes"], 240)
        self.assertEqual(report["summary"][0]["warning_sessions"], 0)
        self.assertEqual(len(report["details"]), 2)

    def test_late_or_early_is_reviewable_not_an_absence(self):
        for records, code, class_id in [(pair("07:05:00"), "LATE", "C1"),
                                        (pair(exited="10:30:00"), "LEFT_EARLY", "C2")]:
            report = self.report(records)
            self.assertEqual(report["summary"][0]["confirmed_sessions"], 1)
            self.assertEqual(report["summary"][0]["warning_sessions"], 1)
            self.assertEqual(report["summary"][0]["recorded_sessions"], 2)
            row = next(row for row in report["details"] if row["class_id"] == class_id)
            self.assertIn(code, row["warnings"])
            self.assertNotIn("MISSED", row["warnings"])

    def test_missing_checkout_never_counts_as_confirmed(self):
        report = self.report(pair()[:1])
        self.assertEqual(report["summary"][0]["confirmed_sessions"], 0)
        self.assertEqual(report["summary"][0]["unverified_events"], 2)

    def test_repeated_checkins_cannot_double_count_a_class(self):
        records = pair()
        records += [{**records[0], "record_id": "in2"}, {**records[1], "record_id": "out2", "check_in_record_id": "in2"}]
        self.assertEqual(self.report(records)["summary"][0]["confirmed_sessions"], 2)

    def test_invalid_location_and_manual_records(self):
        invalid = self.report(pair(status="INVALID"))
        self.assertEqual(invalid["summary"][0]["confirmed_sessions"], 0)
        self.assertIn("INVALID_LOCATION", invalid["details"][0]["warnings"])
        manual = self.report(pair(status="MANUAL"))
        self.assertEqual(manual["summary"][0]["confirmed_sessions"], 2)
        self.assertIn("MANUAL", manual["details"][0]["notes"])

    def test_legacy_records_are_not_reconstructed_from_current_schedule(self):
        records = pair()
        records[0].pop("session_block")
        records[1].pop("check_in_record_id")
        report = self.report(records)
        self.assertEqual(report["summary"][0]["confirmed_sessions"], 0)
        self.assertEqual(report["summary"][0]["unverified_events"], 2)

    def test_inferred_legacy_double(self):
        records = pair("07:49:00", "11:02:00")
        records[0].pop("session_block")
        records[1].pop("check_in_record_id")
        for record in records:
            record["nearest_center_name"] = "Center"
        report = self.report(records)
        self.assertEqual(report["summary"][0]["confirmed_sessions"], 2)
        self.assertEqual(report["summary"][0]["confirmed_minutes"], 0)
        self.assertTrue(all(row["source"] == "LEGACY_DOUBLE" and row["block_size"] == 2 for row in report["details"]))
        self.assertNotIn("session_block", records[0])
        self.assertEqual(report, self.report(records))
        for field, value in [("status", "INVALID"), ("nearest_center_name", "Other")]:
            changed = [dict(r) for r in records]
            changed[1][field] = value
            self.assertEqual(self.report(changed)["summary"][0]["confirmed_sessions"], 0)

    def test_link_cannot_cross_teacher_or_date(self):
        records = pair()
        records[1]["teacher_id"] = "T2"
        report = self.report(records)
        self.assertEqual(report["summary"][0]["confirmed_sessions"], 0)
        records = pair()
        records[1]["checked_at"] = "2026-10-20T11:00:00+07:00"
        self.assertEqual(self.report(records)["summary"][0]["confirmed_sessions"], 0)

    def test_workbook_is_localized_and_has_teacher_detail(self):
        report = self.report(pair())
        for language, title, monday, confirmed in [
            ("vi", "Tổng hợp", "Thứ Hai", "Hợp lệ"),
            ("en", "Summary", "Monday", "Confirmed"),
        ]:
            workbook = load_workbook(build_report_workbook(report, language))
            self.assertEqual(workbook.sheetnames[0], title)
            self.assertEqual(len(workbook.worksheets), 2)
            summary, detail = workbook.worksheets
            self.assertEqual(summary.cell(4, 3).value, 2)
            self.assertEqual(detail.cell(4, 2).value, monday)
            self.assertEqual(detail.cell(4, 9).value, confirmed)
            self.assertEqual(detail.max_row, 5)
            self.assertEqual(detail.freeze_panes, "A4")

    def test_names_are_never_excel_formulas(self):
        teachers = [{**TEACHERS[0], "teacher_name": "=1+1", "teacher_id": "T/1"}]
        report = build_attendance_report(teachers, [], DAY, DAY)
        workbook = load_workbook(build_report_workbook(report, "en"))
        self.assertEqual(workbook.worksheets[0].cell(4, 2).data_type, "s")
        self.assertEqual(workbook.sheetnames, ["Summary"])

    def test_invalid_saved_interval_cannot_be_confirmed(self):
        for end in ("07:00:00", "06:00:00"):
            report = self.report(pair(session_block=[{**BLOCK[0], "end_time": end}]))
            self.assertEqual(report["summary"][0]["confirmed_sessions"], 0)
            self.assertEqual(report["summary"][0]["unverified_events"], 1)
            self.assertIn("INVALID_SHIFT", report["details"][0]["warnings"])

    def test_explicit_link_and_single_saved_shift_are_sufficient_evidence(self):
        records = pair(exited="08:30:00", session_block=None, session_class_id="C1",
                       session_start_time="07:00:00", session_end_time="08:30:00")
        self.assertEqual(self.report(records)["summary"][0]["confirmed_sessions"], 1)

    def test_checkout_uses_saved_double_shift_not_changed_schedule(self):
        db = MagicMock()
        db.execute.return_value.scalar.return_value = "C2"
        check_in = pair()[0]
        resolved = resolve_checkout_session(db, check_in, "T1", "CS1", datetime(2026, 10, 19, 11, tzinfo=timezone(timedelta(hours=7))))
        self.assertEqual(resolved["snapshot"]["class_id"], "C2")
        self.assertEqual(resolved["attendance_status"], "COMPLETED")
        record = snapshot_checkout(db, pair()[1], resolved, check_in)
        self.assertEqual(record["check_in_record_id"], "in1")
        self.assertEqual(record["session_end_time"], "11:00:00")

    def test_checkin_persists_the_whole_contiguous_block(self):
        db = MagicMock()
        db.execute.return_value.mappings.return_value.one.return_value = {
            "session_class_id": "C1", "session_center_id": "CS1",
            "session_start_time": "07:00:00", "session_end_time": "08:30:00",
        }
        db.execute.return_value.mappings.return_value.all.return_value = BLOCK
        record = snapshot_attendance_session(db, pair()[0])
        self.assertEqual(record["session_block"], BLOCK)
        query, params = db.execute.call_args.args
        self.assertIn("session_block", str(query))
        self.assertEqual(json.loads(params["block"]), BLOCK)
        recursive_query = str(db.execute.call_args_list[1].args[0])
        self.assertIn("c.start_time = b.end_time", recursive_query)
        self.assertIn("c.center_id = b.center_id", recursive_query)
        self.assertIn("JOIN assigned_classes", recursive_query)

    def test_report_filter_parameters_are_bound_and_inclusive(self):
        db = MagicMock()
        db.execute.return_value.mappings.return_value.all.side_effect = [TEACHERS, pair()]
        with patch("api.routes.attendance.ensure_attendance_table"):
            report = load_attendance_report(db, DAY, DAY, "T1", "en")
        self.assertEqual(report["summary"][0]["confirmed_sessions"], 2)
        query, params = db.execute.call_args.args
        self.assertEqual(params["teacher_id"], "T1")
        self.assertIn("CAST(:end_date AS DATE) + 1", str(query))
        self.assertIn("ar.teacher_id = :teacher_id", str(query))
        self.assertNotIn("JOIN classes", str(query))

    def test_range_validation_and_unknown_teacher(self):
        for start, end in [(date(2026, 2, 2), date(2026, 2, 1)), (date(2025, 1, 1), date(2026, 1, 2))]:
            with self.assertRaises(HTTPException):
                load_attendance_report(MagicMock(), start, end, None, "vi")
        db = MagicMock()
        db.execute.return_value.mappings.return_value.all.return_value = []
        with patch("api.routes.attendance.ensure_attendance_table"), self.assertRaises(HTTPException) as error:
            load_attendance_report(db, DAY, DAY, "missing", "en")
        self.assertEqual(error.exception.status_code, 404)

    def test_manual_historical_checkout_search_uses_requested_timestamp(self):
        db = MagicMock()
        timestamp = datetime(2026, 10, 19, 4, tzinfo=timezone.utc)
        find_open_check_in(db, "TC00001", timestamp)
        query, params = db.execute.call_args.args
        self.assertEqual(params["checked_at"], timestamp)
        self.assertNotIn("NOW()", str(query))


class VisitReportTests(unittest.TestCase):
    def legacy(self, minutes=193, status="VALID"):
        records = pair("07:57:00", "11:13:00")
        records[0].pop("session_block")
        records[1].pop("check_in_record_id")
        start = datetime.fromisoformat(records[0]["checked_at"])
        records[1]["checked_at"] = (start + timedelta(minutes=minutes)).isoformat()
        records[1]["status"] = status
        for record in records:
            record["nearest_center_name"] = "Center"
        return records

    def test_invalid_checkout_still_pairs_and_counts_two_invalid_classes(self):
        records = self.legacy(status="INVALID")
        records[1]["attendance_status"] = "OVERTIME"
        report = build_attendance_report(TEACHERS, records, DAY, DAY)
        self.assertEqual(report["summary"][0]["invalid_sessions"], 2)
        self.assertEqual(len(report["sessions"]), 1)
        row = report["sessions"][0]
        self.assertEqual(row["block_size"], 2)
        self.assertEqual(row["result"], "INVALID")
        self.assertEqual(set(row["warnings"]), {"INVALID_LOCATION", "OVERTIME"})
        self.assertEqual(row["notes"], [])

    def test_legacy_pair_duration_boundaries(self):
        for minutes, count in [(179, 2), (180, 1), (240, 1), (241, 2)]:
            report = build_attendance_report(TEACHERS, self.legacy(minutes), DAY, DAY)
            self.assertEqual(len(report["sessions"]), count)

    def test_shared_times_appear_once_in_excel_and_summary_still_counts_two(self):
        report = build_attendance_report(TEACHERS, self.legacy(196), DAY, DAY)
        self.assertEqual(report["summary"][0]["confirmed_sessions"], 2)
        for language in ("vi", "en"):
            book = load_workbook(build_report_workbook(report, language))
            sheet = book.worksheets[1]
            self.assertEqual(sheet.max_row, 5)
            self.assertEqual(sheet.cell(4, 6).value, "07:57")
            self.assertEqual(sheet.cell(4, 7).value, "11:13")
            self.assertEqual(sheet.cell(4, 8).value, 2)
            self.assertEqual(sheet.cell(5, 8).value, 2)
            self.assertIsNone(sheet.cell(4, 10).value)

    def test_sheets_split_by_center_and_totals_exclude_incomplete_visits(self):
        records = self.legacy(196)
        records.append({**records[0], "record_id": "missing-out", "class_id": "C3",
                        "checked_at": "2026-10-19T14:00:00+07:00", "nearest_center_name": "Other"})
        report = build_attendance_report(TEACHERS, records, DAY, DAY)
        book = load_workbook(build_report_workbook(report, "vi"))
        self.assertEqual(len(book.worksheets), 3)
        first, second = book.worksheets[1:]
        self.assertIn("Teacher One", first.title)
        self.assertIn("Center", first.title)
        self.assertIn("Other", second.title)
        self.assertEqual(first.cell(first.max_row, 8).value, 2)
        self.assertEqual(second.cell(second.max_row, 8).value, 0)
        self.assertNotIn(str(first.max_row), first.auto_filter.ref.split(":")[-1])
        for sheet in book:
            self.assertNotIn("UTC", sheet.cell(2, 1).value)
        self.assertEqual(first.cell(3, 7).value, "Giờ ra")
        self.assertEqual(first.cell(3, 8).value, "Số ca trong buổi")
        self.assertEqual(first.cell(3, 9).value, "Kết quả")
        self.assertEqual(first.cell(3, 10).value, "Ghi chú")


if __name__ == "__main__":
    unittest.main()
