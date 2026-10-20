import json
import unittest
from contextlib import ExitStack
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from api.routes import attendance
from core.permissions import get_current_user
from db import get_db
from schemas.scheduling import AttendanceCheckInPayload, AttendanceCheckOutPayload, ManualAttendancePayload
from services.attendance_reports import build_attendance_report, build_report_workbook
from test_attendance_reports import BLOCK, DAY, TEACHERS, pair

NAMES = ["Bé An", "Bé Bình"]
LOCATION = {"latitude": 10.77, "longitude": 106.69, "accuracy_meters": 8}
STUDENTS = {"student_count": 2, "student_names": NAMES}
SESSION_STUDENTS = [
    {"class_id": "C1", "student_count": 2, "student_names": NAMES},
    {"class_id": "C2", "student_count": 1, "student_names": ["Bé Chi"]},
]


class AttendanceStudentValidationTests(unittest.TestCase):
    def test_session_lists_require_distinct_classes_and_valid_counts(self):
        payload = AttendanceCheckOutPayload(**LOCATION, session_students=SESSION_STUDENTS)
        self.assertIsNone(payload.student_names)
        self.assertEqual(payload.session_students[1].student_names, ["Bé Chi"])
        for values in [
            {"session_students": []},
            {"session_students": [SESSION_STUDENTS[0], SESSION_STUDENTS[0]]},
            {"session_students": SESSION_STUDENTS, **STUDENTS},
            {"session_students": [{**SESSION_STUDENTS[0], "student_count": 3}]},
            {"session_students": [{**SESSION_STUDENTS[0], "class_id": ""}]},
        ]:
            with self.subTest(values=values), self.assertRaises(ValidationError):
                AttendanceCheckOutPayload(**LOCATION, **values)
        with self.assertRaises(ValidationError):
            ManualAttendancePayload(teacher_id="T1", center_id="CS1", checked_at="2026-10-19T11:00:00+07:00",
                                    session_students=SESSION_STUDENTS)

    def test_checkout_requires_positive_integer_and_exact_nonempty_names(self):
        for values in [
            {}, {"student_count": 1}, {"student_names": NAMES},
            {"student_count": 0, "student_names": []},
            {"student_count": -1, "student_names": []},
            {"student_count": 1.5, "student_names": ["An"]},
            {"student_count": True, "student_names": ["An"]},
            {"student_count": "2", "student_names": NAMES},
            {"student_count": 3, "student_names": NAMES},
            {"student_count": 1, "student_names": ["  "]},
            {"student_count": 1, "student_names": [123]},
            {"student_count": 1, "student_names": ["a" * 151]},
            {"student_count": 1, "student_names": ["An\nBinh"]},
            {"student_count": 1, "student_names": ["An\x00Binh"]},
        ]:
            with self.subTest(values=values), self.assertRaises(ValidationError):
                AttendanceCheckOutPayload(**LOCATION, **values)

    def test_names_are_trimmed_without_deduplicating_children_with_the_same_name(self):
        payload = AttendanceCheckOutPayload(**LOCATION, student_count=2, student_names=["  Bé An ", "Bé An"])
        self.assertEqual(payload.student_names, ["Bé An", "Bé An"])
        self.assertEqual(payload.student_count, 2)

    def test_report_cell_limit_is_enforced_before_names_could_be_truncated(self):
        names = ["a" * 150] * 216 + ["a" * 149, "b"]
        self.assertEqual(len("\n".join(names)), 32767)
        AttendanceCheckOutPayload(**LOCATION, student_count=218, student_names=names)
        with self.assertRaises(ValidationError):
            AttendanceCheckOutPayload(**LOCATION, student_count=218, student_names=names[:-1] + ["bb"])

    def test_checkin_does_not_require_students_and_manual_checkout_does(self):
        AttendanceCheckInPayload(**LOCATION)
        base = {"teacher_id": "T1", "center_id": "CS1", "checked_at": "2026-10-19T11:00:00+07:00"}
        ManualAttendancePayload(**base)
        for values in [{}, {"student_count": 0, "student_names": []}, {"student_count": 3, "student_names": NAMES}]:
            with self.assertRaises(ValidationError):
                ManualAttendancePayload(**base, event_type="CHECK_OUT", **values)
        payload = ManualAttendancePayload(**base, event_type="CHECK_OUT", **STUDENTS)
        self.assertEqual(payload.student_names, NAMES)
        with self.assertRaises(ValidationError):
            ManualAttendancePayload(**base, **STUDENTS)


class AttendanceStudentApiTests(unittest.TestCase):
    def setUp(self):
        self.db = MagicMock()
        self.user = SimpleNamespace(role="teacher", user_code="TC00001", teacher_id="T1", display_name="Teacher One")
        app = FastAPI()
        app.include_router(attendance.router)
        app.dependency_overrides[get_current_user] = lambda: self.user
        app.dependency_overrides[get_db] = lambda: self.db
        self.client = TestClient(app)
        self.addCleanup(self.client.close)

    def mock_checkout(self, stack, double=False):
        for name in ("ensure_attendance_table", "ensure_assignment_tables"):
            stack.enter_context(patch.object(attendance, name))
        check_in = pair()[0]
        if not double:
            check_in = {**check_in, "session_block": BLOCK[1:]}
        stack.enter_context(patch.object(attendance, "find_open_check_in", return_value=check_in))
        stack.enter_context(patch.object(attendance, "resolve_checkout_session", return_value={
            "class_id": "C2", "attendance_status": "COMPLETED", "snapshot": BLOCK[1],
        }))
        stack.enter_context(patch.object(attendance, "get_closest_center", return_value={
            "center_id": "CS1", "status": "VALID", "name": "Center",
        }))

    def assert_saved_students(self):
        calls = [call for call in self.db.execute.call_args_list if "SET student_names" in str(call.args[0])]
        self.assertEqual(len(calls), 1)
        query, params = calls[0].args
        self.assertEqual(json.loads(params["student_names"]), NAMES)
        self.assertEqual(params["record_id"], "out1")
        self.assertIn("event_type = 'CHECK_OUT'", str(query))
        self.db.commit.assert_called_once()

    def test_api_rejects_incomplete_checkout_before_writing(self):
        for values in [{}, {"student_count": 0, "student_names": []}, {"student_count": 3, "student_names": NAMES}]:
            response = self.client.post("/attendance-records/check-out", json={**LOCATION, **values})
            self.assertEqual(response.status_code, 422)
        self.db.execute.assert_not_called()

    def test_automatic_checkout_saves_names_and_returns_them_with_the_link(self):
        self.db.execute.return_value.mappings.return_value.first.side_effect = [None, pair()[1]]
        with ExitStack() as stack:
            self.mock_checkout(stack)
            response = self.client.post("/attendance-records/check-out", json={**LOCATION, **STUDENTS})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["student_names"], NAMES)
        self.assertEqual(response.json()["student_count"], 2)
        self.assertEqual(response.json()["check_in_record_id"], "in1")
        self.assert_saved_students()

    def test_manual_checkout_uses_the_same_persistence(self):
        self.user.role = "manager"
        self.db.execute.return_value.mappings.return_value.first.side_effect = [
            {"teacher_id": "T1", "user_code": "TC00001", "display_name": "Teacher One"},
            {"center_id": "CS1", "name": "Center", "latitude": 10.77, "longitude": 106.69},
            pair()[1],
        ]
        with ExitStack() as stack:
            self.mock_checkout(stack)
            response = self.client.post("/attendance-records/manual", json={
                "teacher_id": "T1", "center_id": "CS1", "checked_at": "2026-10-19T11:00:00+07:00",
                "event_type": "CHECK_OUT", **STUDENTS,
            })
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["student_names"], NAMES)
        self.assert_saved_students()

    def test_double_checkout_saves_independent_lists_with_one_checkout(self):
        for manual in (False, True):
            with self.subTest(manual=manual):
                self.db.reset_mock()
                self.user.role = "manager" if manual else "teacher"
                self.db.execute.return_value.mappings.return_value.first.side_effect = [
                    *([{"teacher_id": "T1", "user_code": "TC00001", "display_name": "Teacher"},
                       {"center_id": "CS1", "name": "Center", "latitude": 10.77, "longitude": 106.69}] if manual else [None]),
                    pair()[1],
                ]
                with ExitStack() as stack:
                    self.mock_checkout(stack, double=True)
                    response = self.client.post("/attendance-records/manual" if manual else "/attendance-records/check-out",
                        json={**({"teacher_id": "T1", "center_id": "CS1", "checked_at": "2026-10-19T11:00:00+07:00",
                                  "event_type": "CHECK_OUT"} if manual else LOCATION), "session_students": SESSION_STUDENTS})
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.json()["session_students"], SESSION_STUDENTS)
                self.assertIsNone(response.json()["student_names"])
                self.assertEqual(response.json()["check_in_record_id"], "in1")
                inserts = [call for call in self.db.execute.call_args_list if "INSERT INTO attendance_records" in str(call.args[0])]
                self.assertEqual(len(inserts), 1)
                save = next(call for call in self.db.execute.call_args_list if "SET student_names" in str(call.args[0]))
                self.assertEqual(json.loads(save.args[1]["session_students"]), SESSION_STUDENTS)
                self.assertIsNone(save.args[1]["student_names"])
                self.db.commit.assert_called_once()

    def test_double_checkout_rejects_shared_missing_or_foreign_class_before_insert(self):
        for values in [STUDENTS, {"session_students": SESSION_STUDENTS[:1]},
                       {"session_students": [SESSION_STUDENTS[0], {**SESSION_STUDENTS[1], "class_id": "OTHER"}]}]:
            with self.subTest(values=values), ExitStack() as stack:
                self.db.reset_mock()
                self.db.execute.return_value.mappings.return_value.first.return_value = None
                self.mock_checkout(stack, double=True)
                response = self.client.post("/attendance-records/check-out", json={**LOCATION, **values})
                self.assertEqual(response.status_code, 400, response.text)
                self.assertTrue(all("INSERT" not in str(call.args[0]) for call in self.db.execute.call_args_list))
                self.db.commit.assert_not_called()

    def test_history_reads_persisted_names_without_exposing_other_teachers(self):
        self.db.execute.return_value.mappings.return_value.all.return_value = [{**pair()[1], **STUDENTS, "session_students": SESSION_STUDENTS}]
        with patch.object(attendance, "ensure_attendance_table"):
            response = self.client.get("/attendance-records")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()[0]["student_names"], NAMES)
        self.assertEqual(response.json()[0]["session_students"], SESSION_STUDENTS)
        query, params = self.db.execute.call_args.args
        self.assertIn("ar.student_names", str(query))
        self.assertIn("ar.session_students", str(query))
        self.assertIn("jsonb_array_length(ar.student_names)", str(query))
        self.assertIn("ar.user_code = :user_code", str(query))
        self.assertEqual(params["user_code"], "TC00001")

    def test_teacher_cannot_submit_manual_records_or_access_reports(self):
        response = self.client.post("/attendance-records/manual", json={
            "teacher_id": "T2", "center_id": "CS1", "checked_at": "2026-10-19T11:00:00+07:00",
            "event_type": "CHECK_OUT", **STUDENTS,
        })
        self.assertEqual(response.status_code, 403)
        for endpoint in ("report", "export"):
            response = self.client.get(f"/attendance-records/{endpoint}?start_date=2026-10-19&end_date=2026-10-19")
            self.assertEqual(response.status_code, 403)
        self.db.execute.assert_not_called()


class AttendanceStudentReportTests(unittest.TestCase):
    def test_split_rosters_stay_separate_in_json_and_excel(self):
        records = pair()
        records[1]["session_students"] = SESSION_STUDENTS
        report = build_attendance_report(TEACHERS, records, DAY, DAY)
        self.assertEqual(len(report["sessions"]), 1)
        self.assertIsNone(report["sessions"][0]["student_names"])
        self.assertEqual([item["student_names"] for item in report["sessions"][0]["session_students"]], [NAMES, ["Bé Chi"]])
        self.assertEqual([item["student_count"] for item in report["details"]], [2, 1])
        self.assertEqual(report["summary"][0]["confirmed_sessions"], 2)
        for language in ("vi", "en"):
            sheet = load_workbook(build_report_workbook(report, language)).worksheets[1]
            self.assertEqual(sheet.cell(4, 4).value, "C1")
            self.assertEqual(sheet.cell(5, 4).value, "C2")
            self.assertEqual(sheet.cell(4, 11).value, 2)
            self.assertEqual(sheet.cell(5, 11).value, 1)
            self.assertEqual(sheet.cell(4, 12).value, "\n".join(NAMES))
            self.assertEqual(sheet.cell(5, 12).value, "Bé Chi")
            self.assertEqual(sheet.cell(6, 8).value, 2)

    def test_double_shift_shares_students_once_per_visit_in_json_and_workbook(self):
        records = pair()
        records[1]["student_names"] = NAMES
        report = build_attendance_report(TEACHERS, records, DAY, DAY)
        self.assertEqual(len(report["sessions"]), 1)
        self.assertEqual(report["sessions"][0]["block_size"], 2)
        self.assertEqual(report["sessions"][0]["student_count"], 2)
        self.assertEqual(report["sessions"][0]["student_names"], NAMES)
        self.assertTrue(all(row["student_names"] == NAMES for row in report["details"]))
        for language, count, names in [
            ("vi", "Số học viên", "Tên học viên"),
            ("en", "Number of students", "Student names"),
        ]:
            sheet = load_workbook(build_report_workbook(report, language)).worksheets[1]
            self.assertEqual(sheet.cell(3, 11).value, count)
            self.assertEqual(sheet.cell(3, 12).value, names)
            self.assertEqual(sheet.cell(4, 11).value, 2)
            self.assertEqual(sheet.cell(4, 12).value, "\n".join(NAMES))
            self.assertEqual(sheet.max_row, 5)

    def test_old_records_remain_unknown_and_do_not_change_session_counts(self):
        report = build_attendance_report(TEACHERS, pair(), DAY, DAY)
        self.assertEqual(report["summary"][0]["confirmed_sessions"], 2)
        self.assertIsNone(report["sessions"][0]["student_count"])
        self.assertIsNone(report["sessions"][0]["student_names"])
        for language, label in [("vi", "Chưa ghi nhận"), ("en", "Not recorded")]:
            sheet = load_workbook(build_report_workbook(report, language)).worksheets[1]
            self.assertEqual(sheet.cell(4, 11).value, label)
            self.assertEqual(sheet.cell(4, 12).value, label)

    def test_names_are_not_excel_formulas(self):
        records = pair()
        records[1]["student_names"] = ["=1+1"]
        report = build_attendance_report(TEACHERS, records, DAY, DAY)
        sheet = load_workbook(build_report_workbook(report, "en")).worksheets[1]
        self.assertEqual(sheet.cell(4, 12).value, "=1+1")
        self.assertEqual(sheet.cell(4, 12).data_type, "s")


if __name__ == "__main__":
    unittest.main()
